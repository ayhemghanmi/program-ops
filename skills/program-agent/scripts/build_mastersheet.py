#!/usr/bin/env python3
"""Build a program mastersheet (Index, Program Plan Gantt, Action Items) in the house format.

Usage:  python3 build_mastersheet.py plan.json "Output Name.xlsx"

plan.json shape (dates as YYYY-MM-DD):
{
  "program": "AD Climate Accelerator",
  "client": "Greenline",            # short name used in approval rows, e.g. "KUEC"
  "operator": "Brinc",
  "start_date": "2026-11-02",        # project start shown at the top; grid starts on that week's Monday
  "phases": [
    {"name": "Phase 1: Program Architecture & Blueprint",
     "items": [
        {"type": "task", "task": "...", "owner": "Brinc", "start": "...", "end": "...", "status": "Not Started"},
        {"type": "approval", "task": "Curriculum sign-off", "start": "...", "end": "..."},
        {"type": "milestone", "task": "Curriculum signed off", "owner": "Brinc", "date": "..."}
     ]}
  ],
  "headline_milestone": {"task": "PROGRAM KICK-OFF: 11 January 2027", "owner": "Brinc & Greenline", "date": "2027-01-11"},
  "index": [{"stage": "Program Architecture & Blueprint", "assets": [{"name": "Program Operating Manual", "link": ""}]}],
  "actions": [{"stage": "1.0", "action": "...", "owner": "Brinc", "due": "2026-11-20", "status": "Not Started"}]
}
Statuses: Not Started, In Progress, Complete, Delayed, Blocked. Approvals: Pending, Approved.
"""
import json, sys, math
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

# ---- house palette (matches the KUEC mastersheet)
NAVY, GREY, BAR, DONE = "2E4057", "D8D8D8", "A6C9EB", "8ED873"
APPR_TAG, APPR_ROW, APPR_TEXT = "FF8C00", "FFF0C8", "7B4F00"
MILE_ROW, MILE_TEXT = "FFE5B4", "8B4513"
MONTH_ALT, TODAY = "DBE9F7", "FFC000"
STATUS_FILL = {"Complete": "D9F2D0", "In Progress": "DDEBF7", "Delayed": "FAE2D5", "Blocked": "F4CCCC",
               "Pending": "FFF0C8", "Approved": "D9F2D0"}
STATUSES = ["Not Started", "In Progress", "Complete", "Delayed", "Blocked"]
FONT = "Arial"

def F(**k): return Font(name=FONT, size=k.pop("size", 10), **k)
def fill(c): return PatternFill("solid", fgColor=c)
thin = Side(style="thin", color="E3E3E3")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center")

def d(s): return date.fromisoformat(s) if isinstance(s, str) and s else None
def monday(x): return x - timedelta(days=x.weekday())

def lines_needed(text, width_chars, bold=False):
    """Approximate wrapped lines for Arial 10 in a column of the given Excel width."""
    if not text: return 1
    per_line = width_chars * (1.12 if bold else 1.28) - 2
    return max(1, math.ceil(len(str(text)) / max(1, per_line)))

def height(lines): return max(16, 13 * lines + 4)

def build(plan, out):
    wb = Workbook()
    client, operator = plan["client"], plan.get("operator", "Operator")
    appr_label = f"{client} Approval"
    # ---------- collect rows
    rows = []
    for pi, ph in enumerate(plan["phases"], start=1):
        rows.append(("phase", f"{pi}.0", ph["name"]))
        n = 0
        for it in ph["items"]:
            t = it.get("type", "task")
            if t == "task":
                n += 1
                rows.append(("task", f"{pi}.{n}", it))
            elif t == "approval":
                rows.append(("approval", "✔ APPR", it))
            elif t == "milestone":
                rows.append(("milestone", "★", it))
    hm = plan.get("headline_milestone")

    dates = []
    for kind, _, it in rows:
        if kind == "phase": continue
        for k in ("start", "end", "date"):
            if d(it.get(k)): dates.append(d(it.get(k)))
    if hm and d(hm.get("date")): dates.append(d(hm["date"]))
    start = d(plan.get("start_date")) or min(dates)
    first_week = monday(min([start] + dates))
    last_week = monday(max(dates)) + timedelta(days=7)
    n_weeks = (last_week - first_week).days // 7 + 1

    # ======================= Program Plan
    ws = wb.active; ws.title = "Program Plan"
    widths = {"A": 1.6, "B": 8.5, "C": 21, "D": 52, "E": 16, "F": 11, "G": 11, "H": 13, "I": 1.2}
    for k, v in widths.items(): ws.column_dimensions[k].width = v
    W0 = 10  # column J
    for i in range(n_weeks): ws.column_dimensions[L(W0 + i)].width = 3.6
    lastc = W0 + n_weeks - 1

    ws["B1"] = plan["program"]; ws["B1"].font = F(size=14, bold=True, color=NAVY)
    ws["C2"] = "Project Start Date"; ws["C2"].font = F(bold=True)
    ws["D2"] = start; ws["D2"].number_format = "dd mmm yyyy"; ws["D2"].font = F(); ws["D2"].alignment = Alignment(horizontal="left")
    ws["E2"] = "Today"; ws["E2"].font = F(bold=True); ws["E2"].alignment = Alignment(horizontal="right")
    ws["F2"] = "=TODAY()"; ws["F2"].number_format = "dd mmm yyyy"; ws["F2"].font = F()
    for r, lab in ((1, "Year"), (2, "Month"), (3, "Week")):
        ws.cell(r, 8, lab).font = F(color="595959"); ws.cell(r, 8).alignment = Alignment(horizontal="right", vertical="bottom")

    weeks = [first_week + timedelta(days=7 * i) for i in range(n_weeks)]
    for i, wk in enumerate(weeks):
        c = ws.cell(3, W0 + i, wk); c.number_format = "dd-mmm"
        c.font = F(size=8); c.alignment = Alignment(horizontal="center", vertical="bottom", text_rotation=90)
    # month and year bands, merged, with alternating fills
    def bands(keyf, row, fmt, alt):
        i = 0; k = 0
        while i < n_weeks:
            j = i
            while j + 1 < n_weeks and keyf(weeks[j + 1]) == keyf(weeks[i]): j += 1
            c = ws.cell(row, W0 + i, date(weeks[i].year, weeks[i].month if fmt == "mmm" else 1, 1))
            c.number_format = fmt; c.font = F(size=9, bold=True); c.alignment = CENTER
            if j > i: ws.merge_cells(start_row=row, start_column=W0 + i, end_row=row, end_column=W0 + j)
            for x in range(i, j + 1):
                ws.cell(row, W0 + x).fill = fill(alt[k % len(alt)])
            i = j + 1; k += 1
    bands(lambda w: (w.year, w.month), 2, "mmm", [MONTH_ALT, "FFFFFF"])
    bands(lambda w: w.year, 1, "yyyy", [GREY, "EDEDED"])
    ws.row_dimensions[3].height = 40

    hdr = ["#", "Stage", "Task", "Responsible", "Start", "End", "Status", ""]
    for j, h in enumerate(hdr, start=2):
        c = ws.cell(4, j, h); c.font = F(bold=True); c.fill = fill(GREY); c.alignment = Alignment(horizontal="left", vertical="center")
    for i in range(n_weeks): ws.cell(4, W0 + i).fill = fill(GREY)
    ws.freeze_panes = "J5"

    r = 5; first_data = r
    status_cells, appr_cells = [], []
    def gantt_cells(rr, bg=None):
        for i in range(n_weeks):
            col = L(W0 + i)
            c = ws.cell(rr, W0 + i, f'=IF(AND($G{rr}<>"",{col}$3=($G{rr}-WEEKDAY($G{rr},2)+1)),"●","")')
            c.font = F(size=7, color="1F3864"); c.alignment = CENTER
            if bg: c.fill = fill(bg)
            c.border = Border(left=thin)

    for kind, num, it in rows:
        if kind == "phase":
            ws.cell(r, 2, num); ws.cell(r, 3, it)
            for col in range(2, lastc + 1):
                ws.cell(r, col).fill = fill(NAVY); ws.cell(r, col).font = F(bold=True, color="FFFFFF")
            ws.cell(r, 3).alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[r].height = 18
        elif kind == "task":
            vals = [num, "", it["task"], it.get("owner", operator), d(it.get("start")), d(it.get("end")), it.get("status", "Not Started")]
            for j, v in enumerate(vals, start=2):
                c = ws.cell(r, j, v); c.font = F(); c.alignment = LEFT
            ws.cell(r, 2).alignment = Alignment(horizontal="left", vertical="center")
            for j in (6, 7): ws.cell(r, j).number_format = "dd-mmm-yy"
            gantt_cells(r); status_cells.append(f"H{r}")
            ws.row_dimensions[r].height = height(lines_needed(it["task"], widths["D"]))
        elif kind == "approval":
            text = f"✔ {client.upper()} APPROVAL: {it['task']}"
            vals = [num, appr_label, text, it.get("owner", client), d(it.get("start")), d(it.get("end")), it.get("status", "Pending")]
            for j, v in enumerate(vals, start=2):
                c = ws.cell(r, j, v); c.fill = fill(APPR_ROW); c.alignment = LEFT
                c.font = F(bold=True, color=APPR_TEXT) if j in (3, 4) else F(bold=j in (5, 8))
            b = ws.cell(r, 2); b.fill = fill(APPR_TAG); b.font = F(size=9, bold=True, color="FFFFFF"); b.alignment = CENTER
            ws.cell(r, 9).fill = fill(APPR_ROW)
            for j in (6, 7): ws.cell(r, j).number_format = "dd-mmm-yy"
            gantt_cells(r, APPR_ROW); appr_cells.append(f"H{r}")
            ws.row_dimensions[r].height = height(lines_needed(text, widths["D"], bold=True))
        elif kind == "milestone":
            dt = d(it.get("date") or it.get("end"))
            text = f"⬛ MILESTONE: {it['task']}"
            vals = [num, "", text, it.get("owner", operator), dt, dt, it.get("status", "Not Started")]
            for j, v in enumerate(vals, start=2):
                c = ws.cell(r, j, v); c.font = F(); c.alignment = LEFT
            for j in (2, 4):
                ws.cell(r, j).fill = fill(MILE_ROW); ws.cell(r, j).font = F(bold=True, color=MILE_TEXT)
            ws.cell(r, 2).alignment = CENTER
            for j in (6, 7): ws.cell(r, j).number_format = "dd-mmm-yy"
            gantt_cells(r); status_cells.append(f"H{r}")
            ws.row_dimensions[r].height = height(lines_needed(text, widths["D"], bold=True))
        r += 1
    if hm:
        dt = d(hm["date"])
        vals = ["★", "", hm["task"], hm.get("owner", operator), dt, dt, hm.get("status", "Not Started")]
        for j, v in enumerate(vals, start=2):
            c = ws.cell(r, j, v); c.font = F(); c.alignment = LEFT
        for j in (2, 3, 4):
            ws.cell(r, j).fill = fill(NAVY); ws.cell(r, j).font = F(bold=True, color=TODAY)
        ws.cell(r, 2).alignment = CENTER
        for j in (6, 7): ws.cell(r, j).number_format = "dd-mmm-yy"
        gantt_cells(r); status_cells.append(f"H{r}")
        ws.row_dimensions[r].height = 18
        r += 1
    last_data = r - 1

    # bars, completion, today, status colours
    rng = f"J{first_data}:{L(lastc)}{last_data}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($H{first_data}="Complete",$G{first_data}<>"",J$3=$G{first_data}-WEEKDAY($G{first_data},2)+1)'], fill=fill(DONE), stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($F{first_data}<>"",$G{first_data}<>"",J$3+6>=$F{first_data},J$3<=$G{first_data})'], fill=fill(BAR)))
    ws.conditional_formatting.add(f"J3:{L(lastc)}4", FormulaRule(formula=['AND(TODAY()>=J$3,TODAY()<J$3+7)'], fill=fill(TODAY), font=Font(bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['AND(TODAY()>=J$3,TODAY()<J$3+7)'], border=Border(left=Side(style="medium", color=TODAY), right=Side(style="medium", color=TODAY))))
    hr = f"H{first_data}:H{last_data}"
    for s, col in STATUS_FILL.items():
        ws.conditional_formatting.add(hr, CellIsRule(operator="equal", formula=[f'"{s}"'], fill=fill(col)))

    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"Pending,Approved"', allow_blank=True)
    ws.add_data_validation(dv); ws.add_data_validation(dv2)
    for c in status_cells: dv.add(c)
    for c in appr_cells: dv2.add(c)

    # legend
    r = last_data + 3
    ws.cell(r, 2, "Legend").font = F(bold=True)
    legend = [("Phase", NAVY, "FFFFFF"), ("✔ APPR  Client approval", APPR_TAG, "FFFFFF"), ("★  Milestone", MILE_ROW, MILE_TEXT),
              ("Planned weeks", BAR, "000000"), ("Completed (end week)", DONE, "000000"), ("This week", TODAY, "000000")]
    for i, (lab, bg, fc) in enumerate(legend, start=1):
        c = ws.cell(r + i, 2); c.fill = fill(bg)
        ws.cell(r + i, 3, lab).font = F(size=9, color="000000")
    ws.cell(r, 5, "Status").font = F(bold=True)
    for i, s in enumerate(STATUSES + ["Pending", "Approved"], start=1):
        c = ws.cell(r + i, 5, s); c.font = F(size=9)
        if s in STATUS_FILL: c.fill = fill(STATUS_FILL[s])
    ws.cell(r + 9, 2, "● marks each item's end week. Bars and the end marker update automatically when you change Start, End or Status.").font = F(size=8, italic=True, color="595959")

    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:4"; ws.print_options.horizontalCentered = True
    ws.sheet_view.zoomScale = 90

    # ======================= Index
    wi = wb.create_sheet("Index", 0)
    for k, v in {"A": 1.6, "B": 5, "C": 34, "D": 40, "E": 44, "F": 14}.items(): wi.column_dimensions[k].width = v
    wi["B1"] = f"{plan['program']}: asset register"; wi["B1"].font = F(size=13, bold=True, color=NAVY)
    for j, h in enumerate(["#", "Stage", "Asset", "Link", "Status"], start=2):
        c = wi.cell(3, j, h); c.font = F(bold=True); c.fill = fill(GREY)
    rr = 4; n = 0
    idv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True); wi.add_data_validation(idv)
    for st in plan.get("index", []):
        n += 1
        wi.cell(rr, 2, f"{n}.0").font = F(bold=True, color="FFFFFF")
        wi.cell(rr, 3, st["stage"]).font = F(bold=True, color="FFFFFF")
        for j in range(2, 7): wi.cell(rr, j).fill = fill(NAVY)
        rr += 1; m = 0
        for a in st.get("assets", []):
            m += 1
            wi.cell(rr, 2, f"{n}.{m}").font = F()
            wi.cell(rr, 4, a["name"]).font = F()
            if a.get("link"):
                c = wi.cell(rr, 5, a["link"]); c.hyperlink = a["link"]; c.font = F(color="0563C1", underline="single")
            else:
                wi.cell(rr, 5, "Add link").font = F(italic=True, color="A6A6A6")
            s = wi.cell(rr, 6, a.get("status", "Not Started")); s.font = F(); idv.add(s.coordinate)
            rr += 1
    for s, col in STATUS_FILL.items():
        wi.conditional_formatting.add(f"F4:F{max(rr, 5)}", CellIsRule(operator="equal", formula=[f'"{s}"'], fill=fill(col)))
    wi.freeze_panes = "B4"

    # ======================= Action Items
    wa = wb.create_sheet("Action Items")
    for k, v in {"A": 1.6, "B": 6, "C": 44, "D": 64, "E": 18, "F": 12, "G": 14}.items(): wa.column_dimensions[k].width = v
    wa["B1"] = f"{plan['program']}: action items"; wa["B1"].font = F(size=13, bold=True, color=NAVY)
    for j, h in enumerate(["#", "Stage", "Action Item", "Responsible", "Due", "Status"], start=2):
        c = wa.cell(3, j, h); c.font = F(bold=True); c.fill = fill(GREY)
    adv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True); wa.add_data_validation(adv)
    phase_names = {f"{i}.0": p["name"] for i, p in enumerate(plan["phases"], start=1)}
    rr = 4
    for i, a in enumerate(plan.get("actions", []), start=1):
        vals = [i, phase_names.get(a.get("stage"), a.get("stage", "")), a["action"], a.get("owner", ""), d(a.get("due")), a.get("status", "Not Started")]
        for j, v in enumerate(vals, start=2):
            c = wa.cell(rr, j, v); c.font = F(); c.alignment = LEFT
        wa.cell(rr, 6).number_format = "dd-mmm-yy"; adv.add(f"G{rr}")
        wa.row_dimensions[rr].height = height(max(lines_needed(a["action"], 64), lines_needed(vals[1], 44)))
        rr += 1
    for k in range(rr, rr + 25):
        adv.add(f"G{k}"); wa.cell(k, 6).number_format = "dd-mmm-yy"
    end = rr + 25
    for s, col in STATUS_FILL.items():
        wa.conditional_formatting.add(f"G4:G{end}", CellIsRule(operator="equal", formula=[f'"{s}"'], fill=fill(col)))
    wa.conditional_formatting.add(f"F4:F{end}", FormulaRule(formula=['AND(F4<>"",F4<TODAY(),$G4<>"Complete")'], font=Font(color="C00000", bold=True)))
    wa.freeze_panes = "B4"

    for sh in (wi, wa):
        sh.page_setup.orientation = "landscape"; sh.sheet_properties.pageSetUpPr.fitToPage = True
        sh.page_setup.fitToWidth = 1; sh.page_setup.fitToHeight = 0
    wb.active = wb.sheetnames.index("Program Plan")
    wb.save(out)
    return {"rows": len(rows) + (1 if hm else 0), "weeks": n_weeks, "first_week": str(first_week), "last_week": str(weeks[-1])}

if __name__ == "__main__":
    plan = json.load(open(sys.argv[1]))
    print(build(plan, sys.argv[2]))
